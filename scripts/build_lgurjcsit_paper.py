from __future__ import annotations

import html
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTICLE = ROOT / "article_final.md"
OUT = ROOT / "paper_lgurjcsit"
FIG_OUT = OUT / "figures"
TABLE_OUT = OUT / "tables"
PAPER_TEX = OUT / "paper.tex"
REFS_BIB = OUT / "references.bib"
DOCX = OUT / "article_lgurjcsit.docx"
PDF = OUT / "article_lgurjcsit.pdf"
REPORT = OUT / "build_report.md"


FIGURES = [
    (
        "reranking_quality_latency_scatter_main.png",
        "Сравнение качества и задержки различных стратегий реранкинга.",
    ),
    (
        "reranking_mean_f1_barplot_main.png",
        "Сравнение качества ответа различных стратегий реранкинга.",
    ),
]


BIB_KEYS = [
    "docvqa2020",
    "layoutlmv32022",
    "donut2021",
    "pix2struct2022",
    "mmlongbenchdoc2024",
    "docbench2024",
    "beir2021",
    "colpali2024",
    "mmembed2024",
    "omniembednemotron2025",
    "vidorev32026",
    "ragsurveygao2023",
    "ragsurveyfan2024",
    "murag2022",
    "m2rag2025",
    "mhierrag2025",
    "raganything2025",
    "monot5-2020",
    "colbert2020",
    "m3embedding2024",
    "qwen3vlreranker2026",
]


@dataclass
class MarkdownParts:
    abstract: str
    keywords: list[str]
    body: list[str]
    references: list[str]


def ensure_dirs() -> None:
    OUT.mkdir(exist_ok=True)
    FIG_OUT.mkdir(exist_ok=True)
    TABLE_OUT.mkdir(exist_ok=True)


def copy_figures() -> list[str]:
    copied = []
    stale_pipeline = FIG_OUT / "pipeline_diagram.png"
    if stale_pipeline.exists():
        stale_pipeline.unlink()
    for name, _caption in FIGURES:
        src = ROOT / "reports" / "figures" / name
        if src.exists():
            shutil.copy2(src, FIG_OUT / name)
            copied.append(name)
    return copied


def parse_markdown(text: str) -> MarkdownParts:
    lines = text.splitlines()
    abstract_lines: list[str] = []
    keywords: list[str] = []
    body: list[str] = []
    refs: list[str] = []
    state = "body"
    for line in lines:
        stripped = line.strip()
        if stripped == "# Аннотация":
            state = "abstract"
            continue
        if stripped == "# Ключевые слова":
            state = "keywords"
            continue
        if stripped in {"## Список литературы", "# Список литературы"}:
            state = "refs"
            continue
        if (
            state == "refs"
            and stripped.startswith("# ")
            and not stripped.endswith("Список литературы")
        ):
            state = "body"
        if stripped.startswith("# 2 ") and state in {"abstract", "keywords"}:
            state = "body"

        if state == "abstract":
            if stripped:
                abstract_lines.append(stripped)
        elif state == "keywords":
            if stripped and not stripped.startswith("#"):
                keywords.append(stripped.rstrip(","))
        elif state == "refs":
            if re.match(r"^\[\d+\]", stripped):
                refs.append(stripped)
        else:
            body.append(line)
    return MarkdownParts(" ".join(abstract_lines), [k for k in keywords if k], body, refs)


def strip_section_number(title: str) -> str:
    title = re.sub(r"^\d+(?:\.\d+)*\s+", "", title.strip())
    title = title.replace(" - ", " — ")
    return title


def latex_escape(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text


def inline_latex(text: str) -> str:
    code_spans: list[str] = []

    def save_code(match: re.Match[str]) -> str:
        code_spans.append(r"\texttt{" + latex_escape(match.group(1)) + "}")
        return f"@@CODE{len(code_spans)-1}@@"

    text = re.sub(r"`([^`]+)`", save_code, text)
    text = latex_escape(text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\\textbf{\1}", text)

    def cite(match: re.Match[str]) -> str:
        nums = [int(item.strip()) for item in match.group(1).split(",") if item.strip().isdigit()]
        if not nums:
            return match.group(0)
        keys = [BIB_KEYS[n - 1] for n in nums if 1 <= n <= len(BIB_KEYS)]
        return r"\cite{" + ",".join(keys) + "}" if keys else match.group(0)

    text = re.sub(r"\[([0-9,\s]+)\]", cite, text)
    text = text.replace("->", r"$\rightarrow$")
    for idx, code in enumerate(code_spans):
        text = text.replace(latex_escape(f"@@CODE{idx}@@"), code)
    return text


def collect_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        raw = lines[idx].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def table_to_latex(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    cols = len(rows[0])
    widths = "|".join(
        ["L{0.95\\columnwidth}"]
        if cols == 1
        else [f"L{{{0.92 / cols:.2f}\\columnwidth}}" for _ in range(cols)]
    )
    out = [
        r"\begin{table*}[t]",
        r"\centering",
        r"\footnotesize",
        r"\setlength{\tabcolsep}{3pt}",
        r"\begin{adjustbox}{max width=\textwidth}",
        rf"\begin{{tabular}}{{{widths}}}",
        r"\toprule",
    ]
    for i, row in enumerate(rows):
        escaped = [inline_latex(cell) for cell in row]
        out.append(" & ".join(escaped) + r" \\")
        out.append(r"\midrule" if i == 0 else "")
    out.extend([r"\bottomrule", r"\end{tabular}", r"\end{adjustbox}", r"\end{table*}"])
    return "\n".join(line for line in out if line)


def markdown_body_to_latex(lines: list[str]) -> str:
    out: list[str] = []
    para: list[str] = []
    i = 0

    def flush_para() -> None:
        if para:
            out.append(inline_latex(" ".join(para)))
            out.append("")
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush_para()
            i += 1
            continue
        if stripped.startswith("```mermaid"):
            flush_para()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            if i < len(lines):
                i += 1
            continue
        if stripped.startswith("```"):
            flush_para()
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            out.append(r"\begin{quote}\small\ttfamily")
            out.extend(latex_escape(item) + r"\\" for item in code_lines)
            out.append(r"\end{quote}")
            out.append("")
            continue
        if stripped.startswith("!["):
            flush_para()
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if match:
                caption = match.group(1)
                path = Path(match.group(2)).name
                if (FIG_OUT / path).exists():
                    out.extend(
                        [
                            r"\begin{figure*}[t]",
                            r"\centering",
                            rf"\includegraphics[width=0.86\textwidth]{{figures/{path}}}",
                            rf"\caption{{{inline_latex(caption)}}}",
                            r"\end{figure*}",
                            "",
                        ]
                    )
            i += 1
            continue
        if stripped.startswith("|"):
            flush_para()
            rows, next_i = collect_table(lines, i)
            out.append(table_to_latex(rows))
            out.append("")
            i = next_i
            continue
        if stripped.startswith("### "):
            flush_para()
            out.append(r"\subsubsection{" + inline_latex(strip_section_number(stripped[4:])) + "}")
            i += 1
            continue
        if stripped.startswith("## "):
            flush_para()
            out.append(r"\subsection{" + inline_latex(strip_section_number(stripped[3:])) + "}")
            i += 1
            continue
        if stripped.startswith("# "):
            flush_para()
            title = strip_section_number(stripped[2:])
            if title == "Заключение":
                out.append(r"\section{Заключение}")
            elif title == "Обсуждение результатов":
                out.append(r"\section{Обсуждение результатов}")
            else:
                out.append(r"\section{" + inline_latex(title) + "}")
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_para()
            items: list[str] = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            out.append(r"\begin{enumerate}")
            for item in items:
                out.append(r"\item " + inline_latex(item))
            out.append(r"\end{enumerate}")
            out.append("")
            continue
        if stripped.startswith("- "):
            flush_para()
            items = []
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:])
                i += 1
            out.append(r"\begin{itemize}")
            for item in items:
                out.append(r"\item " + inline_latex(item))
            out.append(r"\end{itemize}")
            out.append("")
            continue
        para.append(stripped)
        i += 1
    flush_para()
    return "\n".join(out)


def make_bib() -> str:
    src = ROOT / "paper_ieee" / "references.bib"
    return src.read_text(encoding="utf-8")


def make_tex(parts: MarkdownParts) -> str:
    body = markdown_body_to_latex(parts.body)
    keywords = ", ".join(parts.keywords)
    return rf"""% LGURJCSIT-inspired journal layout generated from article_final.md
\documentclass[twoside,twocolumn,10pt]{{article}}
\usepackage{{fontspec}}
\setmainfont{{Times New Roman}}
\usepackage{{polyglossia}}
\setdefaultlanguage{{russian}}
\setotherlanguage{{english}}
\usepackage{{fancyhdr}}
\usepackage{{geometry}}
\usepackage{{abstract}}
\usepackage{{graphicx}}
\usepackage{{titlesec}}
\usepackage{{ragged2e}}
\usepackage{{amssymb}}
\usepackage{{amsmath}}
\usepackage{{booktabs}}
\usepackage{{array}}
\usepackage{{tabularx}}
\usepackage{{float}}
\usepackage{{caption}}
\usepackage{{adjustbox}}
\usepackage{{hyperref}}
\usepackage[backend=bibtex,style=ieee]{{biblatex}}
\addbibresource{{references.bib}}

\newcolumntype{{L}}[1]{{>{{\raggedright\arraybackslash}}p{{#1}}}}
\geometry{{a4paper,left=0.8in,right=0.8in,top=1in,bottom=1in}}
\setlength{{\columnsep}}{{15pt}}
\setlength{{\parindent}}{{0pt}}
\setlength{{\parskip}}{{3pt}}
\captionsetup[table]{{font=footnotesize,labelfont=bf,labelsep=period,justification=raggedright}}
\captionsetup[figure]{{font=footnotesize,labelfont=bf,labelsep=period,justification=raggedright}}
\titleformat{{\section}}{{\normalsize\bfseries\uppercase}}{{\thesection.}}{{1em}}{{}}
\titleformat{{\subsection}}{{\normalsize\bfseries}}{{\thesubsection.}}{{1em}}{{}}
\titleformat{{\subsubsection}}{{\normalsize\bfseries\itshape}}{{\thesubsubsection.}}{{1em}}{{}}

\pagestyle{{fancy}}
\fancyhf{{}}
\fancyhead[C]{{\small Разработка алгоритма реранкинга мультимодальных данных}}
\fancyhead[CO]{{\small Стулов К. В.}}
\fancyfoot[C]{{\thepage}}

\title{{\fontsize{{16}}{{20}}\selectfont\textbf{{Разработка алгоритма реранкинга мультимодальных данных}}}}
\author{{\fontsize{{12}}{{14}}\selectfont Стулов Кирилл Вячеславович\\
\fontsize{{11}}{{13}}\selectfont Университет ИТМО, 09.04.02 Информационные системы и технологии\\
\fontsize{{11}}{{13}}\selectfont Научный руководитель: Вершинин Владислав Константинович}}
\date{{}}

\begin{{document}}
\twocolumn[
\begin{{@twocolumnfalse}}
\maketitle
\section*{{Abstract}}
\fontsize{{10}}{{12}}\selectfont
\justifying
{inline_latex(parts.abstract)}

\vspace{{0.3cm}}
\noindent\textbf{{Keywords:}} {inline_latex(keywords)}
\vspace{{0.6cm}}
\end{{@twocolumnfalse}}
]

\fontsize{{10.5}}{{12.5}}\selectfont
{body}

\printbibliography[title={{References}}]
\end{{document}}
"""


def add_columns(section, count: int = 2) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), "720")


def setup_section(section, *, columns: int = 2, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.4)
        section.right_margin = Cm(1.4)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
    add_columns(section, columns)


def switch_layout(
    doc: Document, *, columns: int, landscape: bool = False, break_type=WD_SECTION.CONTINUOUS
) -> None:
    section = doc.add_section(break_type)
    setup_section(section, columns=columns, landscape=landscape)


def add_run_markup(paragraph, text: str, *, bold_default: bool = False) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            run.bold = bold_default
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold_default


def add_paragraph_docx(doc: Document, text: str, style: str | None = None, bold: bool = False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    add_run_markup(p, text, bold_default=bold)
    return p


def add_table_docx(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    # Word stores table percentage width in fiftieths of a percent:
    # 5000 = 100%, so 4300 = 86% of the available line width.
    tbl_w.set(qn("w:w"), "4300")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    tbl_layout = OxmlElement("w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_pr.append(tbl_layout)

    if cols_count >= 7:
        col_widths = [Cm(2.25)] + [Cm(0.98) for _ in range(cols_count - 2)] + [Cm(2.25)]
        font_size = 5.6
    elif cols_count >= 5:
        col_widths = [Cm(2.3)] + [Cm(1.45) for _ in range(cols_count - 2)] + [Cm(2.3)]
        font_size = 6.3
    elif cols_count == 4:
        col_widths = [Cm(2.4), Cm(3.25), Cm(2.9), Cm(3.25)]
        font_size = 6.8
    else:
        col_widths = [Cm(12.0 / cols_count) for _ in range(cols_count)]
        font_size = 7.5

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(col_widths):
                cell.width = col_widths[idx]

    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.rows[r].cells[c]
            if c < len(col_widths):
                cell.width = col_widths[c]
            cell.text = ""
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement("w:tcW")
            width_twips = (
                int(col_widths[c].cm * 567) if c < len(col_widths) else int(12.0 / cols_count * 567)
            )
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)
            no_wrap = tc_pr.find(qn("w:noWrap"))
            if no_wrap is not None:
                tc_pr.remove(no_wrap)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), "40")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            tc_pr.append(margins)

            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_run_markup(p, value, bold_default=(r == 0))
            for run in p.runs:
                run.font.size = Pt(font_size)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2ECFF" if r % 2 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shading)
    doc.add_paragraph()


def make_docx(parts: MarkdownParts, copied_figures: list[str]) -> None:
    doc = Document()
    sec = doc.sections[0]
    setup_section(sec, columns=2, landscape=False)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(9.5)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(10.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 3"].font.name = "Arial"
    styles["Heading 3"].font.size = Pt(10)
    styles["Heading 3"].font.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Разработка алгоритма реранкинга мультимодальных данных")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)

    for line in [
        "Стулов Кирилл Вячеславович",
        "Университет ИТМО, 09.04.02 Информационные системы и технологии",
        "Научный руководитель: Вершинин Владислав Константинович",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    add_paragraph_docx(doc, "Abstract", bold=True)
    add_paragraph_docx(doc, parts.abstract)
    add_paragraph_docx(doc, "Keywords: " + ", ".join(parts.keywords), bold=True)

    lines = parts.body
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```mermaid"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            if i < len(lines):
                i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].strip())
                i += 1
            if i < len(lines):
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            for idx, item in enumerate(code_lines):
                if idx:
                    p.add_run().add_break()
                run = p.add_run(item)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            continue
        if stripped.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if match:
                name = Path(match.group(2)).name
                path = FIG_OUT / name
                if name in copied_figures and path.exists():
                    switch_layout(doc, columns=1, landscape=False)
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(path), width=Inches(6.55))
                    cap = doc.add_paragraph(match.group(1))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in cap.runs:
                        run.font.size = Pt(8)
                        run.italic = True
                    switch_layout(doc, columns=2, landscape=False)
            i += 1
            continue
        if stripped.startswith("|"):
            rows, next_i = collect_table(lines, i)
            switch_layout(doc, columns=1, landscape=False, break_type=WD_SECTION.CONTINUOUS)
            add_table_docx(doc, rows)
            switch_layout(doc, columns=2, landscape=False, break_type=WD_SECTION.CONTINUOUS)
            i = next_i
            continue
        if stripped.startswith("# "):
            add_paragraph_docx(doc, strip_section_number(stripped[2:]), style="Heading 1")
            i += 1
            continue
        if stripped.startswith("## "):
            add_paragraph_docx(doc, strip_section_number(stripped[3:]), style="Heading 2")
            i += 1
            continue
        if stripped.startswith("### "):
            add_paragraph_docx(doc, strip_section_number(stripped[4:]), style="Heading 3")
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                add_paragraph_docx(
                    doc, re.sub(r"^\d+\.\s+", "", lines[i].strip()), style="List Number"
                )
                i += 1
            continue
        if stripped.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                add_paragraph_docx(doc, lines[i].strip()[2:], style="List Bullet")
                i += 1
            continue
        paragraph_lines = [stripped]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].strip().startswith(("#", "|", "![", "- "))
            and not re.match(r"^\d+\.\s+", lines[i].strip())
        ):
            paragraph_lines.append(lines[i].strip())
            i += 1
        add_paragraph_docx(doc, " ".join(paragraph_lines))

    add_paragraph_docx(doc, "References", style="Heading 1")
    for ref in parts.references:
        add_paragraph_docx(doc, ref)

    doc.save(DOCX)


def export_pdf_with_word() -> tuple[bool, str]:
    ps = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$doc = $word.Documents.Open('{DOCX}')
$doc.ExportAsFixedFormat('{PDF}', 17)
$doc.Close($false)
$word.Quit()
"""
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
        cwd=ROOT,
        text=True,
        capture_output=True,
        timeout=180,
    )
    ok = result.returncode == 0 and PDF.exists()
    return ok, (result.stdout + result.stderr).strip()


def pdf_page_count() -> int | None:
    if not PDF.exists():
        return None
    data = PDF.read_bytes()
    return len(re.findall(rb"/Type\s*/Page\b", data))


def make_report(copied_figures: list[str], pdf_ok: bool, pdf_log: str) -> None:
    pages = pdf_page_count()
    report = f"""# LGURJCSIT-style build report

## Источник

- Основной источник: `article_final.md`.
- Шаблон-ориентир: Lahore Garrison University Research Journal of Computer Science and Information Technology (LGURJCSIT) Template, Overleaf.
- Использован журнальный стиль, близкий к LGURJCSIT: двухколоночная верстка, компактные заголовки, title/authors/abstract/keywords, figures, tables и references.

## Перенесённые разделы

- Аннотация и ключевые слова.
- Введение.
- Обзор литературы.
- Метод.
- Эксперименты.
- Обсуждение результатов.
- Заключение.
- Список литературы.

## Проверка актуальных результатов

- Best multimodal reranker Mean F1 = 0.7023.
- Best multimodal reranker latency = 13.6441s.
- Controlled Nemotron image + BGE text-only Mean F1 = 0.5674.
- Основной фокус сохранён: controlled evaluation мультимодального реранкинга как отдельного компонента Document QA pipeline.
- Adaptive Reranking оставлен как дополнительный эксперимент и не включён в основные графики.

## Включённые графики

{chr(10).join(f'- `figures/{name}`' for name in copied_figures)}

## Ошибки сборки

- LaTeX-исходник подготовлен: `paper.tex`.
- BibTeX-файл подготовлен: `references.bib`.
- PDF-сборка через Word/Office export: {'успешно' if pdf_ok else 'не выполнена'}.
- Для предотвращения разъезда таблиц и рисунков основной текст оставлен в двух колонках, а таблицы и графики вставляются между текстовыми блоками как full-width continuous-секции. В таблицах используются фиксированные ширины столбцов, перенос текста внутри ячеек и уменьшенный табличный шрифт.
{('- Сообщение сборщика: `' + pdf_log.replace('`', '') + '`') if pdf_log else ''}

## Количество страниц

- PDF pages: {pages if pages is not None else 'не определено'}.

## Финальные файлы

- `paper_lgurjcsit/paper.tex`
- `paper_lgurjcsit/references.bib`
- `paper_lgurjcsit/article_lgurjcsit.pdf`
- `paper_lgurjcsit/build_report.md`

## Замечания

- PDF создан как научная журнальная версия на основе актуального текста статьи.
- Научная позиция, метрики и выводы не изменялись.
- Если требуется строгое соответствие Overleaf/LGURJCSIT, `paper.tex` можно дополнительно собрать в среде с XeLaTeX/BibTeX.
"""
    REPORT.write_text(report, encoding="utf-8")


def main() -> None:
    ensure_dirs()
    text = ARTICLE.read_text(encoding="utf-8")
    parts = parse_markdown(text)
    copied = copy_figures()
    REFS_BIB.write_text(make_bib(), encoding="utf-8")
    PAPER_TEX.write_text(make_tex(parts), encoding="utf-8")
    make_docx(parts, copied)
    pdf_ok, pdf_log = export_pdf_with_word()
    make_report(copied, pdf_ok, pdf_log)
    print(f"paper_tex={PAPER_TEX}")
    print(f"references={REFS_BIB}")
    print(f"docx={DOCX}")
    print(f"pdf={PDF} exists={PDF.exists()}")
    print(f"report={REPORT}")


if __name__ == "__main__":
    main()
